"""Reporting boundaries and shared-renderer regression checks."""
import importlib.util
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
HAS_REPORTING = all(importlib.util.find_spec(name) for name in ['pptx', 'reportlab', 'fitz'])


@unittest.skipUnless(HAS_REPORTING, 'Install requirements-reporting.txt for reporting checks')
class ReportingTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from tools import build_portfolio_reports, report_rendering
        cls.builder = build_portfolio_reports
        cls.renderer = report_rendering

    def test_shared_renderer_exports_matching_slide_formats(self):
        import fitz
        from pptx import Presentation
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory)
            with patch.object(self.renderer, 'OUT', output), patch.object(self.renderer, 'ASSETS', output):
                self.renderer.register_fonts()
                chart = self.renderer.bar_chart(['a', 'b'], [-.2, .3], 'Contributions', 'Log odds',
                                                'chart', signed=True, compact=True)
                deck = self.renderer.Deck()
                page = deck.page('Matching slide formats', 'Renderer regression check')
                deck.text(page, 'Recorded model contributions', 50, 160, 850)
                deck.image(page, chart, y=205, h=260)
                deck.save()
            with fitz.open(output / 'results_presentation.pdf') as pdf:
                self.assertEqual(len(pdf), 1)
                self.assertIn('Matching slide formats', pdf[0].get_text())
            slides = Presentation(output / 'results_presentation.pptx')
            self.assertEqual(len(slides.slides), 1)
            self.assertIn('Matching slide formats', '\n'.join(
                shape.text for shape in slides.slides[0].shapes if shape.has_text_frame))

    def test_markdown_generation_preserves_curated_documents(self):
        evidence = ROOT / 'docs/evidence/portfolio_v1'
        primary = pd.read_csv(evidence / 'comparison.csv').assign(study='planned')
        followup = pd.read_csv(evidence / 'intercept_comparison.csv').assign(study='adaptive follow-up')
        data = {'all_summary': pd.concat([primary, followup], ignore_index=True),
                'paired': pd.read_csv(evidence / 'paired_ablations.csv'),
                'stability': pd.read_csv(evidence / 'feature_selection_stability.csv')}
        for key, name in [('audit', 'data_audit.json'), ('permutation', 'permutation_control.json'),
                          ('bootstrap', 'conditional_bootstrap.json'),
                          ('interpretation', 'interpretation_summary.json')]:
            data[key] = json.loads((evidence / name).read_text(encoding='utf-8'))
        with tempfile.TemporaryDirectory() as directory:
            repository = Path(directory)
            (repository / 'docs').mkdir()
            output = repository / 'reports/generated'
            output.mkdir(parents=True)
            readme = repository / 'README.md'
            technical = repository / 'docs/portfolio-audit.md'
            readme.write_text('Curated README', encoding='utf-8')
            technical.write_text('Curated technical report', encoding='utf-8')
            with patch.object(self.builder, 'ROOT', repository), patch.object(self.builder, 'REPORT', output):
                text = self.builder.report_markdown(data)
            self.assertEqual(readme.read_text(encoding='utf-8'), 'Curated README')
            self.assertEqual(technical.read_text(encoding='utf-8'), 'Curated technical report')
            self.assertEqual((output / 'technical_report.md').read_text(encoding='utf-8'), text)
            for name in ['selector_glassbox', 'selector_blackbox']:
                self.assertIn(name, text)
            self.assertNotIn('superseded presentation', text)
            self.assertNotIn('published', text.lower())

    def test_report_cli_help_has_no_output_side_effects(self):
        with tempfile.TemporaryDirectory() as directory:
            result = subprocess.run([sys.executable, str(ROOT / 'tools/build_portfolio_reports.py'), '--help'],
                                    cwd=directory, capture_output=True, text=True, check=True)
            self.assertIn('--report-dir', result.stdout)
            self.assertEqual(list(Path(directory).iterdir()), [])


if __name__ == '__main__':
    unittest.main()
