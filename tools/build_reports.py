"""Build local presentation and theory PDFs from saved artifacts; never retrain.

Run from the repository root with .venv/Scripts/python.exe tools/build_reports.py.
Optional dependencies are in requirements-reporting.txt. All outputs stay in reports/.
"""
from __future__ import annotations
import hashlib
import importlib.metadata
import json
from pathlib import Path
import sys
import textwrap
from xml.sax.saxutils import escape

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import joblib
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from scipy.special import expit
from sklearn.metrics import precision_recall_curve, roc_curve, fbeta_score
from src.config import load_config
from src.data import CreditDataLoader
from src.evaluation import classification_metrics
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_LEFT
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image, PageBreak, Table, TableStyle, KeepTogether
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

OUT = ROOT / 'reports'
ASSETS = OUT / 'figures'
DATA = OUT / 'data'
NAVY, TEAL, GOLD, RED, GRAY, PAPER = '#152B3A', '#087F8C', '#B36B00', '#B83B52', '#546774', '#F6F8F9'
MODEL_COLORS = {'glassbox': TEAL, 'blackbox': GOLD}
LABELS = {'glassbox': 'Glassbox / logistic', 'blackbox': 'Blackbox / XGBoost'}


def read_json(path):
    return json.loads(path.read_text(encoding='utf-8'))


def short(feature, length=48):
    value = feature.replace('numeric__', '').replace('categorical__', '').replace('missingindicator_', 'missing: ')
    return value if len(value) <= length else value[:length-1] + '…'


def savefig(fig, name):
    path = ASSETS / f'{name}.png'
    fig.savefig(path, dpi=190, facecolor='white', bbox_inches='tight')
    plt.close(fig)
    return path


def load_results():
    result = {'models': {}, 'metrics': {}, 'predictions': {}, 'explanations': {}, 'tuning': {}}
    inputs = []
    for name in LABELS:
        folder = ROOT / 'artifacts/selected' / name
        result['models'][name] = joblib.load(folder / 'model.joblib')
        result['explanations'][name] = read_json(folder / 'explanations.json')
        result['tuning'][name] = read_json(ROOT / 'artifacts/tuned' / name / 'tuning.json')
        for variant, parent in [('baseline', ROOT / 'artifacts'), ('tuned', ROOT / 'artifacts/selected')]:
            directory = parent / name
            key = (name, variant)
            result['metrics'][key] = read_json(directory / 'metrics.json')
            result['predictions'][key] = {}
            cfg = result['models'][name].config['evaluation'].copy()
            for split in ['train', 'validation', 'test']:
                frame = pd.read_csv(directory / f'predictions_{split}.csv', float_precision='round_trip')
                result['predictions'][key][split] = frame
                cfg['threshold'] = result['metrics'][key][split]['threshold']
                computed = classification_metrics(frame.actual, frame.probability, cfg)
                for metric in ['roc_auc', 'pr_auc', 'precision', 'recall', 'f1', 'brier_score', 'log_loss']:
                    if not np.isclose(computed[metric], result['metrics'][key][split][metric], atol=1e-10):
                        raise ValueError(f'Artifact mismatch: {key} {split} {metric}')
                if computed['confusion_matrix'] != result['metrics'][key][split]['confusion_matrix']:
                    raise ValueError('Confusion matrix does not match exported predictions')
                inputs.append(directory / f'predictions_{split}.csv')
            inputs += [directory / 'metrics.json', directory / 'config.yaml']
        inputs += [folder / 'model.joblib', folder / 'explanations.json',
                   ROOT / 'artifacts/tuned' / name / 'tuning.json', ROOT / 'artifacts/tuned' / name / 'cv_results.csv']
    reference = result['predictions'][('glassbox', 'tuned')]['test'][['row_index', 'actual']]
    for splits in result['predictions'].values():
        pd.testing.assert_frame_equal(reference, splits['test'][['row_index', 'actual']])
    cfg = load_config(ROOT / 'config/tuned_glassbox.yaml')
    loader = CreditDataLoader(cfg)
    result['X'], result['y'] = loader.load()
    result['splits'] = loader.split(result['X'], result['y'])
    for split, (features, labels) in result['splits'].items():
        saved = result['predictions'][('glassbox', 'tuned')][split]
        np.testing.assert_array_equal(features.index, saved.row_index)
        np.testing.assert_array_equal(labels, saved.actual)
    rows = result['explanations']['glassbox']['row_indices']
    assert rows == result['explanations']['blackbox']['row_indices']
    positives = [i for i, row in enumerate(rows) if result['y'].loc[row] == 1]
    available = positives or list(range(len(rows)))
    case = max(available, key=lambda i: result['explanations']['blackbox']['base_default_probability'][i])
    result['case_position'], result['case_row'] = case, rows[case]
    result['source_hashes'] = {str(path.relative_to(ROOT)): hashlib.sha256(path.read_bytes()).hexdigest() for path in inputs}
    result['checks'] = {}
    for name, explanation in result['explanations'].items():
        sample=result['X'].loc[explanation['row_indices']]
        np.testing.assert_allclose(result['models'][name].predict_proba(sample)[:,1],
                                   explanation['predicted_default_probability'],atol=1e-7)
        reconstructed = expit(np.asarray(explanation['shap']['values']).sum(axis=1) + np.asarray(explanation['shap']['base_values']))
        error = float(np.max(np.abs(reconstructed - np.asarray(explanation['base_default_probability']))))
        assert error < 1e-5, (name, error)
        result['checks'][name + '_shap_probability_max_error'] = error
    glass = result['explanations']['glassbox']
    native = expit(np.asarray(glass['native']['local_log_odds_contributions']).sum(axis=1) + glass['native']['intercept'])
    np.testing.assert_allclose(native, glass['base_default_probability'], atol=1e-10)
    return result


def supplementary_lime(result):
    """One report-only case; same configured algorithm, extra fidelity diagnostics."""
    from lime.lime_tabular import LimeTabularExplainer
    output = {}
    for name, model in result['models'].items():
        before = joblib.hash(model)
        cfg = model.config['explainability']
        train = result['splits']['train'][0]
        background = train.sample(min(len(train), cfg['background_size']), random_state=42)
        pre = model.pipeline_.named_steps['preprocessor']
        classifier = model.pipeline_.named_steps['classifier']
        reference = pre.transform(background)
        row = pre.transform(result['X'].loc[[result['case_row']]])[0]
        names = pre.get_feature_names_out().tolist()
        explainer = LimeTabularExplainer(reference, feature_names=names,
            class_names=['non-default', 'default'], mode='classification', random_state=42,
            discretize_continuous=False)
        explanation = explainer.explain_instance(row, classifier.predict_proba, labels=(1,),
            num_features=cfg['top_features'], num_samples=cfg['lime_num_samples'])
        output[name] = dict(row_index=int(result['case_row']), actual=int(result['y'].loc[result['case_row']]),
            model_probability=float(classifier.predict_proba(row.reshape(1, -1))[0, 1]),
            surrogate_local_prediction=float(np.asarray(explanation.local_pred).ravel()[0]),
            weighted_training_r2=float(explanation.score), intercept=float(explanation.intercept[1]),
            weights=[{'feature': feature, 'weight': float(weight)} for feature, weight in explanation.as_list(label=1)],
            random_seed=42, background_rows=len(background), perturbations=cfg['lime_num_samples'],
            features=cfg['top_features'], provenance='Supplementary one-case LIME generated for this report; no retraining',
            scope='Base classifier probability; transformed and LIME-standardized feature space')
        assert joblib.hash(model) == before, 'Explaining changed fitted model state'
    (DATA / 'supplementary_lime.json').write_text(json.dumps(output, indent=2), encoding='utf-8')
    result['lime'] = output


def bar_chart(names, values, title, xlabel, path, signed=False, color=TEAL, top=12, compact=False):
    values = np.asarray(values)
    indices = np.argsort(np.abs(values))[-top:]
    fig, ax = plt.subplots(figsize=(6.0, 4.8) if compact else (9.0, 4.8))
    selected = values[indices]
    ax.barh(range(len(indices)), selected, color=[RED if value > 0 else TEAL for value in selected] if signed else color)
    ax.set_yticks(range(len(indices)), [short(names[i], 31 if compact else 46) for i in indices], fontsize=10)
    ax.axvline(0, color='#A7B0B6', lw=0.7)
    ax.set_xlabel(xlabel, fontsize=9 if compact else 10)
    ax.set_title(title, loc='left', fontsize=12 if compact else 14, weight='bold', pad=14)
    ax.grid(axis='x', alpha=0.15)
    fig.tight_layout()
    return savefig(fig, path)


def make_charts(result):
    plt.rcParams.update({'font.family': 'DejaVu Sans', 'axes.spines.top': False,
        'axes.spines.right': False, 'axes.labelcolor': NAVY, 'text.color': NAVY,
        'xtick.color': GRAY, 'ytick.color': GRAY, 'axes.titleweight': 'bold'})
    charts = {}
    fig, axes = plt.subplots(1, 2, figsize=(11.5, 4.1), gridspec_kw={'width_ratios': [1, 1.6]})
    counts = result['y'].value_counts()
    axes[0].bar(['Non-default', 'Default'], [counts[0], counts[1]], color=[NAVY, RED])
    for x, value in enumerate([counts[0], counts[1]]): axes[0].text(x, value + 55, f'{value:,}', ha='center', weight='bold')
    axes[0].set_ylim(0, 3200); axes[0].set_title('Only 6.1% defaults'); axes[0].set_ylabel('Rows')
    for j, (split, (_, labels)) in enumerate(result['splits'].items()):
        axes[1].barh(j, len(labels) - labels.sum(), color=NAVY)
        axes[1].barh(j, labels.sum(), left=len(labels) - labels.sum(), color=RED)
        axes[1].text(len(labels) + 30, j, f'{len(labels):,} rows / {labels.sum()} defaults', va='center', fontsize=10)
    axes[1].set_yticks(range(3), ['Train', 'Validation', 'Test']); axes[1].invert_yaxis()
    axes[1].set_xlim(0, 2600); axes[1].set_title('Shared stratified partitions'); axes[1].set_xlabel('Rows')
    fig.tight_layout(); charts['data'] = savefig(fig, 'data_and_splits')
    fig, axes = plt.subplots(1, 2, figsize=(11.5, 4.5))
    for name in LABELS:
        for variant in ['baseline', 'tuned']:
            frame = result['predictions'][(name, variant)]['test']
            metric = result['metrics'][(name, variant)]['test']
            style = '--' if variant == 'baseline' else '-'
            fpr, tpr, _ = roc_curve(frame.actual, frame.probability)
            axes[0].plot(fpr, tpr, style, color=MODEL_COLORS[name], label=f'{name} {variant}: {metric["roc_auc"]:.3f}')
            precision, recall, _ = precision_recall_curve(frame.actual, frame.probability)
            axes[1].step(recall, precision, where='post', ls=style, color=MODEL_COLORS[name], label=f'{name} {variant}: AP {metric["pr_auc"]:.3f}')
    axes[0].plot([0, 1], [0, 1], ':', color=GRAY); axes[0].set(xlabel='False-positive rate', ylabel='Recall / true-positive rate', title='ROC: ranking across all thresholds')
    axes[1].axhline(37/600, color=GRAY, ls=':'); axes[1].set(xlabel='Recall', ylabel='Precision', title='PR: useful under class imbalance', ylim=(0, 0.65))
    for ax in axes: ax.legend(fontsize=8, loc='upper right'); ax.grid(alpha=0.15)
    fig.tight_layout(); charts['ranking'] = savefig(fig, 'ranking_curves')
    fig, axes = plt.subplots(1, 4, figsize=(12, 3.3))
    for ax, (name, variant) in zip(axes, [(n, v) for n in LABELS for v in ['baseline', 'tuned']]):
        m = np.asarray(result['metrics'][(name, variant)]['test']['confusion_matrix'])
        ax.imshow(m, cmap='Blues', vmin=0, vmax=600)
        for i in range(2):
            for j in range(2): ax.text(j, i, str(m[i, j]), ha='center', va='center', fontsize=18, weight='bold', color='white' if m[i,j] > 300 else NAVY)
        ax.set_title(f'{name}\n{variant}', fontsize=12); ax.set_xticks([0,1], ['Predict 0','Predict 1'], fontsize=8)
        ax.set_yticks([0,1], ['Actual 0','Actual 1'], fontsize=8)
    fig.tight_layout(); charts['confusion'] = savefig(fig, 'confusion_matrices')
    fig, axes = plt.subplots(1, 2, figsize=(11.5, 4.2))
    for ax, name in zip(axes, LABELS):
        frame = result['predictions'][(name, 'tuned')]['test']
        precision, recall, thresholds = precision_recall_curve(frame.actual, frame.probability)
        ax.plot(thresholds, precision[:-1], color=NAVY, label='Test precision')
        ax.plot(thresholds, recall[:-1], color=MODEL_COLORS[name], label='Test recall')
        t = result['metrics'][(name, 'tuned')]['test']['threshold']
        ax.axvline(t, ls=':', color=RED, label=f'Validation-selected t={t:.3f}')
        ax.set(xlabel='Threshold', ylabel='Rate', ylim=(0, 1.04), title=LABELS[name]); ax.legend(fontsize=8); ax.grid(alpha=0.15)
    fig.tight_layout(); charts['thresholds'] = savefig(fig, 'threshold_tradeoffs')
    fig, axes = plt.subplots(1, 2, figsize=(11.5, 4.3))
    for name in LABELS:
        for variant in ['baseline','tuned']:
            m = result['metrics'][(name, variant)]['test']; c = m['calibration_curve']
            axes[0].plot(c['mean_prediction'], c['default_rate'], marker='o', ms=4,
                ls='--' if variant == 'baseline' else '-', color=MODEL_COLORS[name], label=f'{name} {variant}')
        frame = result['predictions'][(name, 'tuned')]['test']
        axes[1].hist(frame.probability, bins=np.linspace(0,1,21), alpha=0.5, color=MODEL_COLORS[name], label=f'{name} tuned')
    axes[0].plot([0,1],[0,1], ':', color=GRAY); axes[0].set(xlabel='Mean predicted score (quantile bins)', ylabel='Observed default fraction', title='Reliability: scores vs outcomes', xlim=(0,1), ylim=(0,1))
    axes[1].set(xlabel='Predicted default score', ylabel='Test rows', title='Glassbox scores shifted upward')
    for ax in axes: ax.legend(fontsize=8); ax.grid(alpha=0.15)
    fig.tight_layout(); charts['calibration'] = savefig(fig, 'calibration_and_scores')
    records = []
    for name, ex in result['explanations'].items():
        names = ex['feature_names']; shap = np.asarray(ex['shap']['values']); i = result['case_position']
        native = ex['native']['coefficients' if name == 'glassbox' else 'feature_importances']
        for feature, weight, mean_abs in zip(names, native, ex['shap']['mean_absolute_importance']):
            records.append(dict(model=name, feature=feature, native_weight=weight, shap_mean_absolute_log_odds=mean_abs))
        charts[name + '_native'] = bar_chart(names, native,
            'Largest signed logistic coefficients' if name == 'glassbox' else 'Largest normalized mean-gain importances',
            'Coefficient in log-odds units' if name == 'glassbox' else 'Normalized mean split gain; not a direction',
            name + '_native', signed=name == 'glassbox', color=MODEL_COLORS[name])
        charts[name + '_global_shap'] = bar_chart(names, ex['shap']['mean_absolute_importance'],
            name.capitalize() + ' / mean |SHAP|', 'Mean |SHAP| in log odds; 20 test rows', name + '_global_shap', color=MODEL_COLORS[name],top=8,compact=True)
        values = shap[i]; order = np.argsort(np.abs(values))[::-1][:9]
        labels = [short(names[k], 44) for k in order] + [f'Other {len(values)-len(order)} transformed features']
        selected = list(values[order]) + [float(values.sum() - values[order].sum())]
        base = float(np.asarray(ex['shap']['base_values'])[i]); total = base + sum(selected)
        fig, ax = plt.subplots(figsize=(10.8, 4.9)); current = base
        for j, value in enumerate(selected):
            ax.barh(j, abs(value), left=min(current, current + value), height=.65, color=RED if value>=0 else TEAL)
            ax.text(current + value, j, f' {value:+.3f}', va='center', fontsize=8, ha='left' if value>=0 else 'right')
            ax.plot([current+value]*2, [j+.33, j+.67], color=GRAY, lw=.6); current += value
        ax.set_yticks(range(len(labels)), labels, fontsize=8.5); ax.invert_yaxis()
        ax.axvline(base, color=GRAY, ls=':', label=f'Background margin {base:.3f}')
        ax.axvline(total, color=NAVY, ls='--', label=f'Final margin {total:.3f}; score {expit(total):.3f}')
        ax.set_xlabel('Cumulative base-model log odds'); ax.set_title(f'{name.capitalize()} / row {result["case_row"]}', loc='left', fontsize=14)
        ax.legend(fontsize=8, loc='lower right'); ax.grid(axis='x', alpha=.15)
        fig.tight_layout(); charts[name + '_local_shap'] = savefig(fig, name + '_local_shap')
        lime = result['lime'][name]
        charts[name + '_lime'] = bar_chart([x['feature'] for x in lime['weights']], [x['weight'] for x in lime['weights']],
            f'{name.capitalize()} / LIME', 'Local surrogate coefficient', name + '_lime', signed=True, top=8,compact=True)
    pd.DataFrame(records).to_csv(DATA / 'feature_outputs.csv', index=False)
    fig, axes = plt.subplots(1,2,figsize=(11.5,4.1))
    for ax, name in zip(axes,LABELS):
        frame = pd.read_csv(ROOT / 'artifacts/tuned' / name / 'cv_results.csv')
        x = np.arange(len(frame))
        ax.errorbar(x,frame.mean_test_average_precision,yerr=frame.std_test_average_precision,fmt='.',color=MODEL_COLORS[name],alpha=.8,capsize=2)
        ax.axhline(110/1800,ls=':',color=GRAY); ax.set(title=LABELS[name],xlabel='Candidates ordered by CV AP',ylabel='Mean AP ± fold SD')
        ax.grid(alpha=.15)
    fig.tight_layout(); charts['cv'] = savefig(fig,'cv_search_scores')
    return charts


def register_fonts():
    font_dir = Path(matplotlib.get_data_path()) / 'fonts/ttf'
    pdfmetrics.registerFont(TTFont('ReportSans', str(font_dir / 'DejaVuSans.ttf')))
    pdfmetrics.registerFont(TTFont('ReportSans-Bold', str(font_dir / 'DejaVuSans-Bold.ttf')))
    pdfmetrics.registerFontFamily('ReportSans', normal='ReportSans', bold='ReportSans-Bold', italic='ReportSans', boldItalic='ReportSans-Bold')


def wrap(text, width, size, bold=False):
    font = 'ReportSans-Bold' if bold else 'ReportSans'
    lines = []
    for paragraph in str(text).split('\n'):
        line = ''
        for word in paragraph.split():
            if line and pdfmetrics.stringWidth(line+' '+word,font,size)>width:
                lines.append(line); line=word
            else: line=(line+' '+word).strip()
        lines.append(line)
    return lines


class Deck:
    """Shared page geometry for editable PowerPoint text and matching PDF slides."""
    def __init__(self): self.pages=[]
    def page(self,title,subtitle='',source='Saved project artifacts',dark=False):
        page={'title':title,'subtitle':subtitle,'source':source,'dark':dark,'objects':[]}; self.pages.append(page)
        return page
    def box(self,p,x,y,w,h,color): p['objects'].append(('box',x,y,w,h,color))
    def text(self,p,text,x,y,w,size=18,color=NAVY,bold=False,leading=None):
        lines=wrap(text,w,size,bold); leading=leading or size*1.35
        p['objects'].append(('text',lines,x,y,w,size,color,bold,leading))
        if y + len(lines)*leading > 493: raise ValueError(f'Slide text exceeds body: {p["title"]}: {text[:60]}')
        return y+len(lines)*leading
    def image(self,p,path,x=40,y=145,w=880,h=325): p['objects'].append(('image',path,x,y,w,h))
    def bullets(self,p,items,x=50,y=150,w=850,size=20,gap=18):
        for item in items:
            self.box(p,x,y+7,4,9,TEAL)
            y=self.text(p,item,x+18,y,w-18,size=size)+gap
    def cards(self,p,cards,y=155):
        width=(860-20*(len(cards)-1))/len(cards)
        for i,(label,value,detail,color) in enumerate(cards):
            x=50+i*(width+20); self.box(p,x,y,width,128,PAPER)
            self.text(p,label.upper(),x+15,y+14,width-30,11,GRAY,True)
            self.text(p,value,x+15,y+38,width-30,31,color,True)
            self.text(p,detail,x+15,y+89,width-30,12,GRAY)
    def table(self,p,headers,rows,widths=None,x=45,y=150,w=870,rowh=42,size=14):
        widths=widths or [1/len(headers)]*len(headers)
        xpos=np.r_[0,np.cumsum(widths)]*w+x
        self.box(p,x,y,w,rowh,NAVY)
        for i,label in enumerate(headers): self.text(p,label,float(xpos[i])+10,y+10,float(widths[i]*w)-20,size,'#FFFFFF',True)
        for j,row in enumerate(rows):
            yy=y+(j+1)*rowh; self.box(p,x,yy,w,rowh,PAPER if j%2==0 else '#FFFFFF')
            for i,label in enumerate(row): self.text(p,str(label),float(xpos[i])+10,yy+10,float(widths[i]*w)-20,size,NAVY,i==0)
    def save(self):
        prs=Presentation(); prs.slide_width=Pt(960); prs.slide_height=Pt(540)
        pdf=canvas.Canvas(str(OUT/'results_presentation.pdf'),pagesize=(960,540))
        pdf.setTitle('Credit model results and interpretability'); pdf.setAuthor('Local project report')
        from PIL import Image as PILImage
        for index,p in enumerate(self.pages,1):
            slide=prs.slides.add_slide(prs.slide_layouts[6]); bg=NAVY if p['dark'] else '#FFFFFF'
            slide.background.fill.solid(); slide.background.fill.fore_color.rgb=RGBColor.from_string(bg[1:])
            pdf.setFillColor(colors.HexColor(bg)); pdf.rect(0,0,960,540,fill=1,stroke=0)
            foreground='#FFFFFF' if p['dark'] else NAVY
            objects=[('text',['CREDIT MODEL STUDY  /  16 OCT 2026'],45,24,870,10,TEAL if not p['dark'] else '#63D0D0',True,14),
                     ('text',wrap(p['title'],870,29,True),45,51,870,29,foreground,True,35),
                     ('text',wrap(p['subtitle'],870,13),45,122,870,13,'#CFDCE4' if p['dark'] else GRAY,False,17)] + p['objects']
            objects += [('text',wrap(p['source'],810,8),45,512,810,8,'#CFDCE4' if p['dark'] else GRAY,False,10),
                        ('text',[f'{index:02d} / {len(self.pages):02d}'],875,512,55,8,foreground,False,10)]
            for obj in objects:
                if obj[0]=='box':
                    _,x,y,w,h,color=obj; shape=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Pt(x),Pt(y),Pt(w),Pt(h))
                    shape.fill.solid(); shape.fill.fore_color.rgb=RGBColor.from_string(color[1:]); shape.line.fill.background()
                    pdf.setFillColor(colors.HexColor(color)); pdf.rect(x,540-y-h,w,h,fill=1,stroke=0)
                elif obj[0]=='text':
                    _,lines,x,y,w,size,color,bold,leading=obj
                    for line_index,line in enumerate(lines):
                        yy=y+line_index*leading
                        tf=slide.shapes.add_textbox(Pt(x),Pt(yy),Pt(w+3),Pt(leading+3)).text_frame
                        tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0; tf.word_wrap=False
                        paragraph=tf.paragraphs[0]; paragraph.text=line
                        paragraph.font.name='DejaVu Sans'; paragraph.font.size=Pt(size); paragraph.font.bold=bold
                        paragraph.font.color.rgb=RGBColor.from_string(color[1:]); paragraph.space_after=Pt(0)
                        pdf.setFont('ReportSans-Bold' if bold else 'ReportSans',size); pdf.setFillColor(colors.HexColor(color)); pdf.drawString(x,540-yy-size*.98,line)
                else:
                    _,path,x,y,w,h=obj
                    with PILImage.open(path) as im: ratio=im.width/im.height
                    iw=min(w,h*ratio); ih=iw/ratio; ix=x+(w-iw)/2; iy=y+(h-ih)/2
                    slide.shapes.add_picture(str(path),Pt(ix),Pt(iy),width=Pt(iw),height=Pt(ih))
                    pdf.drawImage(str(path),ix,540-iy-ih,width=iw,height=ih,mask='auto')
            slide.notes_slide.notes_text_frame.text=p['title']+'\n'+p['subtitle']+'\nSource: '+p['source']+'\nLocal project artifact; see theory_guide.pdf for definitions and limitations.'
            pdf.showPage()
        prs.save(OUT/'results_presentation.pptx'); pdf.save()


def main():
    for directory in [OUT,ASSETS,DATA]: directory.mkdir(parents=True,exist_ok=True)
    register_fonts()
    result=load_results(); supplementary_lime(result); charts=make_charts(result)
    summary=dict(case_row=int(result['case_row']), case_actual=int(result['y'].loc[result['case_row']]),
        checks=result['checks'], top_features={name:ex['native']['top_features'][:5] for name,ex in result['explanations'].items()},
        lime=result['lime'], source_hashes=result['source_hashes'])
    (DATA/'report_summary.json').write_text(json.dumps(summary,indent=2),encoding='utf-8')
    from report_content import make_deck, make_theory
    deck=make_deck(Deck(),result,charts); deck.save(); make_theory(result,charts,OUT)
    versions={p:importlib.metadata.version(p) for p in ['python-pptx','reportlab','PyMuPDF','matplotlib','numpy','pandas','scikit-learn','shap','lime','xgboost']}
    (DATA/'manifest.json').write_text(json.dumps(dict(source_sha256=result['source_hashes'],versions=versions,
        note='Local-only reports; saved models loaded without fitting; LIME is supplementary'),indent=2),encoding='utf-8')
    (OUT/'README.md').write_text('''# Local results package

- [Editable presentation](results_presentation.pptx): 26 slides covering both models, pros/cons, baseline/tuned results, native explanations, SHAP and supplementary LIME.
- [Presentation PDF](results_presentation.pdf): matching widescreen slides.
- [Theory guide](theory_guide.pdf): 28 chapters, worked examples, explanation figures, glossary and primary-source bibliography.
- `data/`: exact feature outputs, supplementary LIME diagnostics, source hashes, environment versions, extracted PDF text and validation checks.
- `figures/` and `previews/`: local chart assets and rendered review images.

All outputs are local and this directory is ignored by Git. No models were retrained or modified to generate the reports. Native and SHAP outputs come from the saved selected runs. LIME was generated as a labeled supplementary one-case demonstration, using saved classifiers and training background rows. The original baseline is included for comparison.

Regenerate from the repository root:

```powershell
.\\.venv\\Scripts\\python.exe -m pip install -r requirements-reporting.txt
.\\.venv\\Scripts\\python.exe tools/build_reports.py
.\\.venv\\Scripts\\python.exe tools/verify_reports.py
```

The source narrative is `tools/theory_content.md`; slide content is in `tools/report_content.py`. Metric consistency, test-row alignment, model/explanation alignment, SHAP/native reconstruction, PDF text bounds and PowerPoint shape bounds are checked. PDF pages were visually reviewed. The matching slide PDF is generated directly from shared layout instructions; it is not an export from an installed PowerPoint application.
''',encoding='utf-8')
    print(json.dumps(dict(slides=len(deck.pages),case_row=summary['case_row'],checks=result['checks'],outputs=str(OUT)),indent=2))


if __name__=='__main__': main()
