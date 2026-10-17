"""Validate the new local PDFs/PPTX and render review contact sheets only there."""
from pathlib import Path
import json
import argparse
import fitz
from PIL import Image,ImageDraw
from pptx import Presentation
ROOT=Path(__file__).resolve().parents[1]
parser=argparse.ArgumentParser(description=__doc__);parser.add_argument('--report-dir',type=Path,default=ROOT/'reports/generated')
folder=parser.parse_args().report_dir.resolve()
preview=folder/'previews';preview.mkdir(exist_ok=True)
results={}
for name in ['results_presentation.pdf','theory_guide.pdf','technical_report.pdf']:
    doc=fitz.open(folder/name);outside=[];word_count=0
    for number,page in enumerate(doc,1):
        text=page.get_text();word_count+=len(text.split())
        assert len(text.strip())>60,(name,number,'empty page')
        assert '\ufffd' not in text,(name,number,'replacement glyph')
        for block in page.get_text('dict')['blocks']:
            if block['type']!=0:continue
            for line in block['lines']:
                for span in line['spans']:
                    bounds=fitz.Rect(span['bbox'])
                    if bounds.x0<0 or bounds.y0<0 or bounds.x1>page.rect.width+1 or bounds.y1>page.rect.height+1:
                        outside.append(dict(page=number,text=span['text'],bounds=list(bounds)))
    assert not outside,(name,outside)
    results[name]=dict(pages=len(doc),words=word_count,text_bounds='passed')
    cols=4;width=390 if name.startswith('results') else 275
    for start in range(0,len(doc),16):
        numbers=list(range(start,min(start+16,len(doc))))
        height=int(width*doc[0].rect.height/doc[0].rect.width)
        sheet=Image.new('RGB',(cols*(width+12),((len(numbers)+cols-1)//cols)*(height+30)),'#DCE4E9');draw=ImageDraw.Draw(sheet)
        for k,n in enumerate(numbers):
            page=doc[n];pix=page.get_pixmap(matrix=fitz.Matrix(width/page.rect.width,width/page.rect.width),alpha=False)
            im=Image.frombytes('RGB',[pix.width,pix.height],pix.samples)
            x=k%cols*(width+12)+6;y=k//cols*(height+30)+20;sheet.paste(im,(x,y));draw.text((x,y-15),str(n+1),fill='black')
        sheet.save(preview/f'{Path(name).stem}_contact_{start//16+1}.png')
    examples=[8,13,14,17,27,29] if name.startswith('results') else [0,2,min(10,len(doc)-1)]
    for n in examples:
        doc[n].get_pixmap(matrix=fitz.Matrix(1.4,1.4),alpha=False).save(preview/f'{Path(name).stem}_{n+1:02d}.png')
prs=Presentation(folder/'results_presentation.pptx')
assert len(prs.slides)==results['results_presentation.pdf']['pages']
for slide in prs.slides:
    for shape in slide.shapes:
        assert shape.left>=0 and shape.top>=0 and shape.left+shape.width<=prs.slide_width and shape.top+shape.height<=prs.slide_height
results['results_presentation.pptx']=dict(slides=len(prs.slides),shape_bounds='passed')
(folder/'validation.json').write_text(json.dumps(results,indent=2),encoding='utf-8')
print(json.dumps(results,indent=2))
