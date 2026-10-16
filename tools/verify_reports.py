"""Check report structure and render local contact sheets for visual review."""
from pathlib import Path
import json
import fitz
from PIL import Image, ImageDraw
from pptx import Presentation

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'reports'
PREVIEW=OUT/'previews'
PREVIEW.mkdir(exist_ok=True)
summary={}
for filename in ['results_presentation.pdf','theory_guide.pdf']:
    doc=fitz.open(OUT/filename)
    texts=[page.get_text() for page in doc]
    assert all(len(text.strip())>70 for text in texts), 'Empty or nearly empty page'
    assert not any('\ufffd' in text for text in texts), 'Replacement glyph in PDF'
    outside=[]
    for n,page in enumerate(doc,1):
        for block in page.get_text('dict')['blocks']:
            if block['type']!=0: continue
            for line in block['lines']:
                for span in line['spans']:
                    rect=fitz.Rect(span['bbox'])
                    if rect.x0<0 or rect.y0<0 or rect.x1>page.rect.width+1 or rect.y1>page.rect.height+1:
                        outside.append({'page':n,'text':span['text'],'bounds':list(rect)})
    assert not outside, outside
    summary[filename]={'pages':len(doc),'words':sum(len(t.split()) for t in texts),'text_outside_page':outside}
    (OUT/'data'/filename.replace('.pdf','_text.txt')).write_text('\n\n'.join(texts),encoding='utf-8')
    per_sheet=16 if filename.startswith('theory') else 28
    for batch in range(0,len(doc),per_sheet):
        pages=list(range(batch,min(len(doc),batch+per_sheet)))
        columns=4; thumb_width=460 if filename.startswith('results') else 300
        ratio=doc[0].rect.height/doc[0].rect.width
        thumb_height=int(thumb_width*ratio)
        sheet=Image.new('RGB',(columns*(thumb_width+12),((len(pages)+columns-1)//columns)*(thumb_height+32)), '#DEE5E9')
        draw=ImageDraw.Draw(sheet)
        for j,n in enumerate(pages):
            pix=doc[n].get_pixmap(matrix=fitz.Matrix(thumb_width/doc[n].rect.width,thumb_width/doc[n].rect.width),alpha=False)
            im=Image.frombytes('RGB',[pix.width,pix.height],pix.samples)
            x=(j%columns)*(thumb_width+12)+6; y=(j//columns)*(thumb_height+32)+20
            sheet.paste(im,(x,y)); draw.text((x,y-16),str(n+1),fill='black')
        sheet.save(PREVIEW/f'{Path(filename).stem}_contact_{batch//per_sheet+1}.png')
    if filename.startswith('results'):
        for n in [7,11,14,17,18,19,21]:
            doc[n].get_pixmap(matrix=fitz.Matrix(1.5,1.5),alpha=False).save(PREVIEW/f'slide_{n+1:02d}.png')
    else:
        for n in [0,1,2,23,25]:
            doc[n].get_pixmap(matrix=fitz.Matrix(1.4,1.4),alpha=False).save(PREVIEW/f'theory_{n+1:02d}.png')

prs=Presentation(OUT/'results_presentation.pptx')
assert len(prs.slides)==26
for i,slide in enumerate(prs.slides,1):
    for shape in slide.shapes:
        assert shape.left>=0 and shape.top>=0
        assert shape.left+shape.width<=prs.slide_width
        assert shape.top+shape.height<=prs.slide_height
summary['results_presentation.pptx']={'slides':len(prs.slides),'shape_bounds':'passed'}
summary['lime']={name:{key:values[key] for key in ['row_index','model_probability','surrogate_local_prediction','weighted_training_r2']}
                 for name,values in json.loads((OUT/'data/supplementary_lime.json').read_text()).items()}
(OUT/'data/validation.json').write_text(json.dumps(summary,indent=2),encoding='utf-8')
print(json.dumps(summary,indent=2))
