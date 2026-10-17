"""Chart helpers and shared PDF/PowerPoint geometry for portfolio reports."""
from pathlib import Path
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib import colors
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

OUT = Path(__file__).resolve().parents[1] / 'reports/generated'
ASSETS = OUT / 'figures'
NAVY, TEAL, GOLD, RED, GRAY, PAPER = '#152B3A', '#087F8C', '#B36B00', '#B83B52', '#546774', '#F6F8F9'

def short(feature, length=48):
    value = feature.replace('numeric__', '').replace('categorical__', '').replace('missingindicator_', 'missing: ')
    return value if len(value) <= length else value[:length-1] + '…'

def savefig(fig, name):
    path = ASSETS / f'{name}.png'
    fig.savefig(path, dpi=190, facecolor='white', bbox_inches='tight')
    plt.close(fig)
    return path

def bar_chart(names, values, title, xlabel, path, signed=False, color=TEAL, top=12, compact=False):
    values = np.asarray(values)
    indices = np.argsort(np.abs(values))[-top:]
    fig, ax = plt.subplots(figsize=(6.0, 4.8) if compact else (9.0, 4.8))
    selected = values[indices]
    ax.barh(range(len(indices)), selected, color=[RED if value > 0 else TEAL for value in selected] if signed else color)
    ax.set_yticks(range(len(indices)), [short(names[i], 31 if compact else 46) for i in indices], fontsize=10)
    ax.axvline(0, color='#A7B0B6', lw=0.7)
    ax.set_xlabel(xlabel, fontsize=9 if compact else 10)
    ax.set_title(title, loc='left', fontsize=12 if compact else 14, weight='bold', pad=14)
    ax.grid(axis='x', alpha=0.15)
    fig.tight_layout()
    return savefig(fig, path)

def register_fonts():
    font_dir = Path(matplotlib.get_data_path()) / 'fonts/ttf'
    pdfmetrics.registerFont(TTFont('ReportSans', str(font_dir / 'DejaVuSans.ttf')))
    pdfmetrics.registerFont(TTFont('ReportSans-Bold', str(font_dir / 'DejaVuSans-Bold.ttf')))
    pdfmetrics.registerFontFamily('ReportSans', normal='ReportSans', bold='ReportSans-Bold', italic='ReportSans', boldItalic='ReportSans-Bold')

def wrap(text, width, size, bold=False):
    font = 'ReportSans-Bold' if bold else 'ReportSans'
    lines = []
    for paragraph in str(text).split('\n'):
        line = ''
        for word in paragraph.split():
            if line and pdfmetrics.stringWidth(line+' '+word,font,size)>width:
                lines.append(line); line=word
            else: line=(line+' '+word).strip()
        lines.append(line)
    return lines

class Deck:
    """Shared page geometry for editable PowerPoint text and matching PDF slides."""
    def __init__(self): self.pages=[]
    def page(self,title,subtitle='',source='Saved project artifacts',dark=False):
        page={'title':title,'subtitle':subtitle,'source':source,'dark':dark,'objects':[]}; self.pages.append(page)
        return page
    def box(self,p,x,y,w,h,color): p['objects'].append(('box',x,y,w,h,color))
    def text(self,p,text,x,y,w,size=18,color=NAVY,bold=False,leading=None):
        lines=wrap(text,w,size,bold); leading=leading or size*1.35
        p['objects'].append(('text',lines,x,y,w,size,color,bold,leading))
        if y + len(lines)*leading > 493: raise ValueError(f'Slide text exceeds body: {p["title"]}: {text[:60]}')
        return y+len(lines)*leading
    def image(self,p,path,x=40,y=145,w=880,h=325): p['objects'].append(('image',path,x,y,w,h))
    def bullets(self,p,items,x=50,y=150,w=850,size=20,gap=18):
        for item in items:
            self.box(p,x,y+7,4,9,TEAL)
            y=self.text(p,item,x+18,y,w-18,size=size)+gap
    def cards(self,p,cards,y=155):
        width=(860-20*(len(cards)-1))/len(cards)
        for i,(label,value,detail,color) in enumerate(cards):
            x=50+i*(width+20); self.box(p,x,y,width,128,PAPER)
            self.text(p,label.upper(),x+15,y+14,width-30,11,GRAY,True)
            self.text(p,value,x+15,y+38,width-30,31,color,True)
            self.text(p,detail,x+15,y+89,width-30,12,GRAY)
    def table(self,p,headers,rows,widths=None,x=45,y=150,w=870,rowh=42,size=14):
        widths=widths or [1/len(headers)]*len(headers)
        xpos=np.r_[0,np.cumsum(widths)]*w+x
        self.box(p,x,y,w,rowh,NAVY)
        for i,label in enumerate(headers): self.text(p,label,float(xpos[i])+10,y+10,float(widths[i]*w)-20,size,'#FFFFFF',True)
        for j,row in enumerate(rows):
            yy=y+(j+1)*rowh; self.box(p,x,yy,w,rowh,PAPER if j%2==0 else '#FFFFFF')
            for i,label in enumerate(row): self.text(p,str(label),float(xpos[i])+10,yy+10,float(widths[i]*w)-20,size,NAVY,i==0)
    def save(self):
        prs=Presentation(); prs.slide_width=Pt(960); prs.slide_height=Pt(540)
        pdf=canvas.Canvas(str(OUT/'results_presentation.pdf'),pagesize=(960,540))
        pdf.setTitle('Credit model results and interpretability'); pdf.setAuthor('Credit-model research')
        from PIL import Image as PILImage
        for index,p in enumerate(self.pages,1):
            slide=prs.slides.add_slide(prs.slide_layouts[6]); bg=NAVY if p['dark'] else '#FFFFFF'
            slide.background.fill.solid(); slide.background.fill.fore_color.rgb=RGBColor.from_string(bg[1:])
            pdf.setFillColor(colors.HexColor(bg)); pdf.rect(0,0,960,540,fill=1,stroke=0)
            foreground='#FFFFFF' if p['dark'] else NAVY
            objects=[('text',['CREDIT MODEL STUDY  /  16 OCT 2026'],45,24,870,10,TEAL if not p['dark'] else '#63D0D0',True,14),
                     ('text',wrap(p['title'],870,29,True),45,51,870,29,foreground,True,35),
                     ('text',wrap(p['subtitle'],870,13),45,122,870,13,'#CFDCE4' if p['dark'] else GRAY,False,17)] + p['objects']
            objects += [('text',wrap(p['source'],810,8),45,512,810,8,'#CFDCE4' if p['dark'] else GRAY,False,10),
                        ('text',[f'{index:02d} / {len(self.pages):02d}'],875,512,55,8,foreground,False,10)]
            for obj in objects:
                if obj[0]=='box':
                    _,x,y,w,h,color=obj; shape=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Pt(x),Pt(y),Pt(w),Pt(h))
                    shape.fill.solid(); shape.fill.fore_color.rgb=RGBColor.from_string(color[1:]); shape.line.fill.background()
                    pdf.setFillColor(colors.HexColor(color)); pdf.rect(x,540-y-h,w,h,fill=1,stroke=0)
                elif obj[0]=='text':
                    _,lines,x,y,w,size,color,bold,leading=obj
                    for line_index,line in enumerate(lines):
                        yy=y+line_index*leading
                        tf=slide.shapes.add_textbox(Pt(x),Pt(yy),Pt(w+3),Pt(leading+3)).text_frame
                        tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0; tf.word_wrap=False
                        paragraph=tf.paragraphs[0]; paragraph.text=line
                        paragraph.font.name='DejaVu Sans'; paragraph.font.size=Pt(size); paragraph.font.bold=bold
                        paragraph.font.color.rgb=RGBColor.from_string(color[1:]); paragraph.space_after=Pt(0)
                        pdf.setFont('ReportSans-Bold' if bold else 'ReportSans',size); pdf.setFillColor(colors.HexColor(color)); pdf.drawString(x,540-yy-size*.98,line)
                else:
                    _,path,x,y,w,h=obj
                    with PILImage.open(path) as im: ratio=im.width/im.height
                    iw=min(w,h*ratio); ih=iw/ratio; ix=x+(w-iw)/2; iy=y+(h-ih)/2
                    slide.shapes.add_picture(str(path),Pt(ix),Pt(iy),width=Pt(iw),height=Pt(ih))
                    pdf.drawImage(str(path),ix,540-iy-ih,width=iw,height=ih,mask='auto')
            slide.notes_slide.notes_text_frame.text=p['title']+'\n'+p['subtitle']+'\nSource: '+p['source']+'\nSee theory_guide.pdf for definitions and interpretation limits.'
            pdf.showPage()
        prs.save(OUT/'results_presentation.pptx'); pdf.save()
